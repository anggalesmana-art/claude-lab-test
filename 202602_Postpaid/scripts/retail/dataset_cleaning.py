import pandas as pd 
import sys 
from openpyxl import load_workbook


def main(file_name):
    print(f"Received input: {file_name}") 
    
    # df = pd.read_excel(f'/data/rclone/sourcepostpaid/source_ipos/{file_name}', 
    #            engine="openpyxl", dtype='str')  
    # Load the workbook and first sheet
    wb = load_workbook(f"/data/gcp/ipos/{file_name}", data_only=True)
    sheet = wb.worksheets[0]
    
    # Get the maximum range used
    max_row = sheet.max_row
    max_col = sheet.max_column
    
    # Extract all cell values
    data = []
    for row in sheet.iter_rows(min_row=1, max_row=max_row, max_col=max_col, values_only=True):
        data.append(list(row))
    
    # Convert to DataFrame (optional: detect headers)
    df = pd.DataFrame(data[1:], columns=data[0])

    # Replace '|' with an empty string in all columns 
    
    df = df.applymap(lambda x: x.replace('|', '') if isinstance(x, str) else x) 
    num_rows = df.shape[0]
    if num_rows == 0:
        print("Excel kosong!")
    elif num_rows == 1:
        print("Excel hanya memiliki 1 baris:", df.iloc[0])
    else:
        print("Excel memiliki lebih dari 1 baris") 
        df.to_excel(f'/data/gcp/ipos/{file_name}', index=False)